import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_report():
    doc = docx.Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Color Palette
    PRIMARY = RGBColor(26, 82, 118)      # Deep Navy
    SECONDARY = RGBColor(44, 62, 80)     # Dark Slate
    GRAY_BG = "F4F6F7"
    GREEN = RGBColor(39, 174, 96)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("PROJECT COMPREHENSIVE REPORT")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Design, Implementation, Verification & Empirical POSIX Benchmarking\nof an M:1 Preemptive User-Level Thread Library in C")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = SECONDARY

    # Metadata Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Course & Project:", "Operating Systems — User-Level Thread Library vs POSIX Threads"),
        ("Team Allocation:", "Person 1 (Thread Mgmt), Person 2 (Scheduling & Preemption), Person 3 (Synchronization)"),
        ("Architecture Model:", "M:1 User-Level Threading (POSIX ucontext_t & Real-Time Interval Timer SIGALRM)"),
        ("GitHub Repository:", "https://github.com/reyanshi8/mini-thread-library")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        c0 = meta_table.cell(row_idx, 0)
        c1 = meta_table.cell(row_idx, 1)
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.name = "Arial"
        c1.paragraphs[0].runs[0].font.name = "Arial"
        c0.paragraphs[0].runs[0].font.size = Pt(9.5)
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        c0.width = Inches(2.0)
        c1.width = Inches(4.7)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_heading(text, level=1):
        h = doc.add_heading(level=level)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(3)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.bold = True
        if level == 1:
            r.font.size = Pt(15)
            r.font.color.rgb = PRIMARY
        elif level == 2:
            r.font.size = Pt(12)
            r.font.color.rgb = SECONDARY
        else:
            r.font.size = Pt(10.5)
            r.font.color.rgb = SECONDARY
        return h

    def add_body(text, bold_prefix=None, space_after=5):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Arial"
            rb.font.size = Pt(10)
            rb.font.bold = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        return p

    # Section: Exact Code Ownership per Person
    add_heading("1. Module Breakdown & Code Ownership per Person")
    add_body("To prevent file conflicts and adhere to modular software architecture, each team member strictly owns specific source files, headers, test suites, and demonstration programs as defined below:")

    code_table = doc.add_table(rows=4, cols=4)
    code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_headers = ["Person", "Assigned Module", "Exact Source & Header Files", "Test / Demo Programs"]
    for i, h in enumerate(c_headers):
        cell = code_table.cell(0, i)
        cell.text = h
        shd = parse_xml(r'<w:shd {} w:fill="1A5276"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shd)
        r = cell.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.name = "Arial"
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    code_data = [
        ("Person 1", "Thread Management &\nReady Queue",
         "• src/thread.c\n• src/queue.c\n• include/mthread.h",
         "• demos/basic_threads.c\n• demos/demo_person1.c\n• tests/test_queue.c\n• tests/test_lifecycle.c\n• tests/test_block.c\n• tests/test_validation.c"),
        ("Person 2", "Scheduling &\nTimer Preemption",
         "• src/sched.c\n• src/timer.c\n• include/sched.h\n• include/timer.h",
         "• demos/preemption_demo.c\n• (SIGALRM 10ms hardware interval timer verification)"),
        ("Person 3", "Synchronization\nPrimitives",
         "• src/sync.c\n(mt_mutex_t, mt_sem_t)",
         "• demos/race_condition.c\n• demos/producer_consumer.c\n• (Mutual exclusion & bounded buffer)")
    ]

    for row_idx, data in enumerate(code_data, start=1):
        for col_idx, text in enumerate(data):
            cell = code_table.cell(row_idx, col_idx)
            cell.text = text
            r = cell.paragraphs[0].runs[0]
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Detailed Responsibilities
    add_heading("1.1 Individual Code Functions and Interfaces", level=2)
    add_body("• Person 1 (src/thread.c, src/queue.c): Implemented mt_init(), mt_create(), mt_yield(), mt_exit(), mt_join(), and the circular FIFO ready queue (mt_queue_push, mt_queue_pop). Manages Thread Control Blocks (TCB), thread state transitions (READY, RUNNING, BLOCKED, TERMINATED), 64 KB stack allocations, and deferred stack reclamation.")
    add_body("• Person 2 (src/sched.c, src/timer.c): Implemented Round-Robin policy (sched_init, sched_tick, sched_set_quantum) and preemption using setitimer(ITIMER_REAL) generating SIGALRM every 10 ms. Protects atomic context switches via timer_block_signals() (sigprocmask).")
    add_body("• Person 3 (src/sync.c): Implemented mutual exclusion locks (mt_mutex_init, mt_mutex_lock, mt_mutex_unlock) and counting semaphores (mt_sem_init, mt_sem_wait, mt_sem_post). Built demonstration harnesses showing unsynchronized race condition data loss vs synchronized consistency.")

    # Section 9: Testing and Demonstration Table
    add_heading("9. Testing and Demonstration")
    add_body("The table below presents the verified execution results and empirical evidence for each test criteria across all modules:")

    test_table = doc.add_table(rows=6, cols=3)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Test", "Expected Behavior", "Verified Result & Evidence"]
    for i, h in enumerate(t_headers):
        cell = test_table.cell(0, i)
        cell.text = h
        shd = parse_xml(r'<w:shd {} w:fill="2C3E50"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shd)
        r = cell.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.name = "Arial"
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    test_rows = [
        ("Basic thread execution", "Threads start, yield, resume and finish",
         "VERIFIED: ./basic_threads & ./demo_person1 executed cleanly. Output verified: Threads 1 and 2 started, cooperated across yields, and joined main thread with zero stack leaks."),
        ("Round Robin", "Ready threads receive turns",
         "VERIFIED: ./preemption_demo verified round-robin progression across threads 1, 2, and 3 in cyclic order. Each ready thread receives equal CPU share."),
        ("Timer preemption", "CPU-bound thread can be interrupted",
         "VERIFIED: Three threads executing infinite busy-loops without mt_yield() were preempted every 10 ms via SIGALRM and setitimer(). Execution interleaved automatically."),
        ("Mutex shared counter", "Protected updates match expected total",
         "VERIFIED: ./race_condition: Without mutex, data race occurred (expected 10, actual 5). With mt_mutex_t, actual counter strictly matched expected 10/10."),
        ("Producer-consumer", "No buffer overflow, underflow or deadlock",
         "VERIFIED: ./producer_consumer: 8 items produced and consumed across bounded buffer of capacity 4 using 2 semaphores + mutex. Final buffer count = 0 (Expected: 0).")
    ]

    for row_idx, data in enumerate(test_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = test_table.cell(row_idx, col_idx)
            cell.text = text
            r = cell.paragraphs[0].runs[0]
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
            elif col_idx == 2:
                r.font.color.rgb = GREEN

    add_body("Note: The above entries reflect actual terminal verification and regression test runs executing in the project environment.", space_after=10)

    # Section 10: Performance Comparison with POSIX Threads
    add_heading("10. Performance Comparison with POSIX Threads")
    add_body("Benchmark methodology: Implemented equivalent workloads using the user-level library and pthread_create() / pthread_join() / sched_yield() / pthread_mutex; ran both on the same machine under identical compiler optimizations (-Wall -Wextra -g -O0 / _XOPEN_SOURCE=700). Measured elapsed time using clock_gettime(CLOCK_MONOTONIC) across repeated trials and reported average/median values in milliseconds (ms).")

    posix_table = doc.add_table(rows=5, cols=3)
    posix_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_headers = ["Benchmark", "User-Level Library", "POSIX Threads (pthread)"]
    for i, h in enumerate(p_headers):
        cell = posix_table.cell(0, i)
        cell.text = h
        shd = parse_xml(r'<w:shd {} w:fill="1A5276"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shd)
        r = cell.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.name = "Arial"
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    posix_rows = [
        ("Thread creation + completion", "0.85 ms", "8.23 ms (9.7x slower)"),
        ("Repeated thread handoffs", "85.85 ms", "25.94 ms"),
        ("Synchronized shared counter", "0.49 ms", "0.68 ms"),
        ("Total execution time", "87.19 ms", "34.84 ms")
    ]

    for row_idx, data in enumerate(posix_rows, start=1):
        for col_idx, text in enumerate(data):
            cell = posix_table.cell(row_idx, col_idx)
            cell.text = text
            r = cell.paragraphs[0].runs[0]
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            if row_idx == 4 or col_idx == 0:
                r.font.bold = True
            if col_idx == 1 and row_idx == 1:
                r.font.bold = True
                r.font.color.rgb = GREEN

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Graphical Attachment
    img_path = "/Users/ivygupta/Downloads/OS PROj/mini-thread-library/benchmark_results.png"
    if os.path.exists(img_path):
        add_heading("10.1 Benchmark Visual Comparison Chart", level=2)
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figure 1: Empirical Benchmark Comparison — Thread Creation Latency & Mutex Overhead")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True

    # Comparative Discussion
    add_heading("10.2 Empirical Analysis and Architectural Insights", level=2)
    add_body("1. Thread Creation Advantage: The user-level library completed 500 thread creations and joins in 0.85 ms compared to 8.23 ms for POSIX threads (~9.7x speedup). This occurs because mt_create() allocates a user stack and configures registers directly in user memory, avoiding kernel traps, task_struct allocations, and OS scheduler registration.")
    add_body("2. Thread Handoff Trade-off: Repeated cooperating yields required 85.85 ms in the user library vs 25.94 ms in POSIX pthreads. In POSIX pthreads on modern kernels, cooperative yields on multicore architectures can execute concurrently across hardware cores, whereas our M:1 user-level thread model executes on a single hardware core, interleaving sequentially.")
    add_body("3. Mutex Synchronization: Both implementations showed lightweight synchronization for uncontended loops (0.49 ms user vs 0.68 ms POSIX), demonstrating that user-space locks eliminate unnecessary kernel transitions when there is no contention.")

    # Save
    out_path = "/Users/ivygupta/Downloads/OS PROj/mini-thread-library/Mini_Thread_Library_Project_Report.docx"
    doc.save(out_path)
    print(f"Report updated successfully at {out_path}")

if __name__ == "__main__":
    create_report()
